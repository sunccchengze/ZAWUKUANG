#!/usr/bin/env python3
"""Build a full-bleed 16:9 image-only PowerPoint from ordered PNG/JPEG pages.

Usage: python scripts/build_image_deck.py OUT.pptx IMG1.png IMG2.png ...
Originals are embedded unaltered; using add_picture(width,height) causes at most
~0.8% aspect adjustment for the 1376x768 generated pages on a 16:9 canvas.
"""
import sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


def main():
    if len(sys.argv) < 4:
        raise SystemExit('Usage: build_image_deck.py OUT.pptx IMG1 IMG2 ...')
    out = Path(sys.argv[1])
    image_paths = [Path(x) for x in sys.argv[2:]]
    prs = Presentation()
    prs.slide_width = Inches(13.3333333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for index, image_path in enumerate(image_paths, start=1):
        if not image_path.is_file():
            raise FileNotFoundError(image_path)
        with Image.open(image_path) as im:
            ratio = im.width / im.height
            if abs(ratio - 16/9) > .03:
                raise ValueError(f'Page {index}: unexpected aspect ratio: {image_path} {im.size}')
            print(f'P{index:02d}  {im.width}x{im.height}  {image_path}')
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(image_path), 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f'Saved {out} ({len(prs.slides)} slides)')


if __name__ == '__main__':
    main()
